"""
Головне вікно програми
"""
from PySide6.QtWidgets import (
    QMainWindow, QWidget, QVBoxLayout, QHBoxLayout,
    QPushButton, QLabel, QProgressBar, QStatusBar,
    QMessageBox, QFileDialog, QInputDialog, QProgressDialog
)
from PySide6.QtCore import Qt, QThread, Signal
import os
import json
from docx import Document

from gui.selection_dialog import SelectionDialog
from gui.add_data_dialog import AddDataDialog
from gui.import_data_dialog import ImportDataDialog
from gui.add_servicemember_dialog import AddServicememberDialog
from gui.add_period_dialog import AddPeriodDialog
from gui.edit_periods_dialog import EditPeriodsDialog
from gui.passport_data_dialog import PassportDataDialog
from core.excel_reader import ExcelReader
from core.database import DatabaseManager
from core.migration import DataMigration
from core.data_processor import DataProcessor
from core.report_generator import ReportGenerator
from utils.validators import validate_excel_file
from utils.paths import get_base_dir, get_resources_dir, get_config_path, get_template_path, get_database_path, get_output_dir


class ReportGeneratorThread(QThread):
    """
    Потік для генерації рапортів у фоновому режимі
    АДАПТОВАНО: працює з ExcelReader або DatabaseManager
    ВИПРАВЛЕНО: Створює власне підключення до БД для уникнення SQLite threading issues
    """
    progress = Signal(int, int)  # (поточний, всього)
    finished = Signal(int, int, list)  # (успішно, помилок, список_помилок)
    error = Signal(str)

    def __init__(self, data_source, names, sheet_names, template_path, output_dir, manual_data=None, report_type="", use_database=False, db_path=None, passport_data_source=None):
        super().__init__()
        self.data_source = data_source  # ExcelReader або None (якщо БД)
        self.use_database = use_database  # Чи використовувати БД
        self.db_path = db_path  # Шлях до БД (для створення підключення в потоці)
        self.names = names
        self.sheet_names = sheet_names
        self.template_path = template_path
        self.output_dir = output_dir
        self.manual_data = manual_data or {}
        self.report_type = report_type
        self.passport_data_source = passport_data_source  # PassportDataDialog або None

    @staticmethod
    def get_initials(full_name: str) -> str:
        """
        Витягує ініціали з ПІБ
        Наприклад: "БАРТ ВОЛОДИМИР ГРИГОРОВИЧ" -> "В.Г."
        """
        parts = full_name.strip().split()
        if len(parts) >= 3:
            # Беремо перші літери імені та по-батькові
            return f"{parts[1][0]}.{parts[2][0]}."
        elif len(parts) == 2:
            return f"{parts[1][0]}."
        return ""

    @staticmethod
    def get_surname(full_name: str) -> str:
        """
        Витягує прізвище з ПІБ
        Наприклад: "БАРТ ВОЛОДИМИР ГРИГОРОВИЧ" -> "БАРТ"
        """
        parts = full_name.strip().split()
        return parts[0] if parts else ""

    @staticmethod
    def get_unit_designation(unit: str) -> str:
        """
        Повертає позначення підрозділу для назви файлу

        Args:
            unit: Назва підрозділу (наприклад, "Г-3")

        Returns:
            Позначення підрозділу (наприклад, "12шр_4шб")
        """
        unit_map = {
            "Г-1": "10шр_4шб",
            "Г-2": "11шр_4шб",
            "Г-3": "12шр_4шб",
            "Г-4": "1мб_4шб",
            "Г-5": "2мб_4шб",
            "Г-6": "ГРВ_4шб",
            "Г-7": "ПТВ_4шб",
            "Г-8": "ВПРК_4шб",
            "Г-9": "кв_4шб",
            "Г-10": "рбак_4шб",
            "Г-11": "ЗРВ_4шб",
            "Г-12": "РВ_4шб",
            "Г-13": "ІСВ_4шб",
            "Г-14": "вРЕБ_4шб",
            "Г-15": "ВЗ_4шб",
            "Г-16": "ВТЗ_4шб",
            "Г-17": "ВМЗ_4шб",
            "Г-18": "мп_4шб",
            "Г": "упр_4шб",
            "Ь": "в_розпорядженні"
        }

        # Повертаємо позначення з мапи, або "4шб" якщо підрозділ не знайдено
        return unit_map.get(unit, "4шб")

    def run(self):
        db_manager = None
        try:
            success_count = 0
            error_count = 0
            errors_list = []  # Список помилок
            total = len(self.names)

            # Створюємо папку output якщо її немає
            os.makedirs(self.output_dir, exist_ok=True)

            # ВИПРАВЛЕННЯ: Створюємо DatabaseManager всередині потоку (SQLite threading fix)
            if self.use_database and self.db_path:
                db_manager = DatabaseManager(self.db_path)
                db_manager.connect()
                data_source = db_manager
            else:
                data_source = self.data_source

            generator = ReportGenerator(self.template_path)

            for i, name in enumerate(self.names):
                try:
                    # Агрегуємо дані (АДАПТОВАНО: працює з Excel або БД)
                    data = DataProcessor.aggregate_servicemember_data(
                        data_source,
                        name,
                        self.sheet_names
                    )

                    if data:
                        # Генеруємо рапорт з новим форматом назви
                        surname = self.get_surname(name)
                        initials = self.get_initials(name)
                        unit = data.get("unit", "")
                        unit_designation = self.get_unit_designation(unit)
                        filename = f"Рапорт_{surname}_{initials}_{unit_designation} ({self.report_type}).docx"
                        output_path = os.path.join(self.output_dir, filename)

                        # Отримуємо паспортні дані для цього імені
                        current_manual_data = dict(self.manual_data)  # Копіюємо базові дані
                        if self.passport_data_source:
                            passport_data = self.passport_data_source.get_passport_for_name(name)
                            current_manual_data.update(passport_data)

                        # Генеруємо рапорт (може кинути exception з детальною помилкою)
                        generator.generate_report(data, output_path, current_manual_data)
                        success_count += 1
                    else:
                        error_count += 1
                        errors_list.append(f"{name}: Дані не знайдено")
                except Exception as e:
                    error_count += 1
                    errors_list.append(f"{name}: {str(e)}")

                # Оновлюємо прогрес
                self.progress.emit(i + 1, total)

            self.finished.emit(success_count, error_count, errors_list)

        except Exception as e:
            self.error.emit(str(e))
        finally:
            # Закриваємо підключення до БД якщо воно було створене
            if db_manager:
                db_manager.close()


class MainWindow(QMainWindow):
    """
    Головне вікно додатку
    """

    def __init__(self):
        super().__init__()
        self.excel_reader = None
        self.db_manager = None  # НОВЕ: менеджер БД
        self.use_database = True  # НОВЕ: використовувати БД як primary джерело
        self.config = self.load_config()
        self.init_ui()

    def load_config(self):
        """
        Завантаження конфігурації
        """
        config_path = get_config_path()

        if os.path.exists(config_path):
            with open(config_path, 'r', encoding='utf-8') as f:
                return json.load(f)
        return {}

    def init_ui(self):
        """
        Ініціалізація інтерфейсу
        """
        self.setWindowTitle(self.config.get("ui", {}).get("window_title", "Генератор періодів участі"))
        self.setMinimumSize(
            self.config.get("ui", {}).get("window_width", 800),
            self.config.get("ui", {}).get("window_height", 600)
        )

        # Центральний віджет
        central_widget = QWidget()
        self.setCentralWidget(central_widget)

        layout = QVBoxLayout()

        # Заголовок
        title = QLabel("Генератор вільного часу")
        title.setObjectName("title")
        title.setAlignment(Qt.AlignCenter)
        layout.addWidget(title)

        # Підзаголовок
        subtitle = QLabel("Автоматизація обліку періодів служби")
        subtitle.setAlignment(Qt.AlignCenter)
        subtitle.setStyleSheet("font-size: 12px; color: gray;")
        layout.addWidget(subtitle)

        layout.addSpacing(30)

        # Кнопки
        button_layout = QVBoxLayout()
        button_layout.setSpacing(20)

        self.btn_periods_100 = QPushButton("Періоди на 100 тис.")
        self.btn_periods_100.clicked.connect(self.on_periods_100_clicked)
        button_layout.addWidget(self.btn_periods_100)

        self.btn_pilgova = QPushButton("Пільгова вислуга (100 тис.+30 тис.)")
        self.btn_pilgova.clicked.connect(self.on_pilgova_clicked)
        button_layout.addWidget(self.btn_pilgova)

        self.btn_import_month = QPushButton("Додати новий місяць")
        self.btn_import_month.clicked.connect(self.on_import_month_clicked)
        button_layout.addWidget(self.btn_import_month)

        self.btn_recalculate = QPushButton("Перерахувати періоди")
        self.btn_recalculate.clicked.connect(self.on_recalculate_periods_clicked)
        self.btn_recalculate.setStyleSheet("background-color: #ff9800; color: white;")
        button_layout.addWidget(self.btn_recalculate)

        # Нові кнопки для управління даними
        self.btn_add_servicemember = QPushButton("Додати військовослужбовця")
        self.btn_add_servicemember.clicked.connect(self.on_add_servicemember_clicked)
        self.btn_add_servicemember.setStyleSheet("background-color: #2196f3; color: white;")
        button_layout.addWidget(self.btn_add_servicemember)

        self.btn_add_period = QPushButton("Додати період")
        self.btn_add_period.clicked.connect(self.on_add_period_clicked)
        self.btn_add_period.setStyleSheet("background-color: #4caf50; color: white;")
        button_layout.addWidget(self.btn_add_period)

        self.btn_edit_periods = QPushButton("Редагувати періоди")
        self.btn_edit_periods.clicked.connect(self.on_edit_periods_clicked)
        self.btn_edit_periods.setStyleSheet("background-color: #9c27b0; color: white;")
        button_layout.addWidget(self.btn_edit_periods)

        self.btn_add_data = QPushButton("Імпорт з Excel (опціонально)")
        self.btn_add_data.clicked.connect(self.on_add_data_clicked)
        button_layout.addWidget(self.btn_add_data)

        self.btn_settings = QPushButton("Налаштування")
        self.btn_settings.clicked.connect(self.on_settings_clicked)
        button_layout.addWidget(self.btn_settings)

        self.btn_update = QPushButton("Оновлення")
        self.btn_update.clicked.connect(self.on_update_clicked)
        button_layout.addWidget(self.btn_update)

        layout.addLayout(button_layout)

        layout.addSpacing(30)

        # Прогрес-бар
        self.progress_bar = QProgressBar()
        self.progress_bar.setVisible(False)
        layout.addWidget(self.progress_bar)

        layout.addStretch()

        central_widget.setLayout(layout)

        # Статус-бар
        self.status_bar = QStatusBar()
        self.setStatusBar(self.status_bar)
        self.status_bar.showMessage("Готовий до роботи")

        # Завантаження Excel при запуску
        self.load_excel_file()

    def load_excel_file(self):
        """
        Завантаження Excel файлу та підключення до БД
        ОНОВЛЕНО: Спочатку перевіряємо БД, Excel потрібен тільки якщо БД порожня
        """
        try:
            # Вимикаємо кнопки під час ініціалізації
            self.status_bar.showMessage("Ініціалізація...")
            self.btn_periods_100.setEnabled(False)
            self.btn_pilgova.setEnabled(False)
            self.btn_import_month.setEnabled(False)
            self.btn_recalculate.setEnabled(False)
            self.btn_add_servicemember.setEnabled(False)
            self.btn_add_period.setEnabled(False)
            self.btn_edit_periods.setEnabled(False)
            self.btn_add_data.setEnabled(False)
            self.btn_settings.setEnabled(False)

            # Примушуємо оновити GUI
            from PySide6.QtWidgets import QApplication
            QApplication.processEvents()

            # 1. ПРІОРИТЕТ: Спочатку підключаємося до БД
            db_config = self.config.get("database", {})
            db_path = db_config.get("database_path", "data.db")

            # Отримуємо абсолютний шлях до БД
            db_absolute_path = get_database_path(db_path)

            self.status_bar.showMessage("Підключення до БД...")
            QApplication.processEvents()

            self.db_manager = DatabaseManager(db_absolute_path)
            self.db_manager.connect()

            # Перевірити чи використовувати БД
            self.use_database = db_config.get("use_database_primary", True)

            # 2. Якщо БД НЕ порожня - працюємо ТІЛЬКИ з БД, Excel НЕ потрібен!
            if not self.db_manager.is_empty():
                self.excel_reader = None
                self.status_bar.showMessage("БД готова до роботи")
                QApplication.processEvents()

                # Увімкнюємо кнопки генерації та імпорту
                self.btn_periods_100.setEnabled(True)
                self.btn_pilgova.setEnabled(True)
                self.btn_settings.setEnabled(True)
                self.btn_import_month.setEnabled(True)  # УВІМКНЕНО для БД!
                self.btn_recalculate.setEnabled(True)  # УВІМКНЕНО для БД!

                # НОВІ кнопки управління даними - УВІМКНЕНІ для БД!
                self.btn_add_servicemember.setEnabled(True)
                self.btn_add_period.setEnabled(True)
                self.btn_edit_periods.setEnabled(True)

                # Excel імпорт залишається опціональним
                self.btn_add_data.setEnabled(True)

                self.status_bar.showMessage("Готово до роботи (джерело: БД)")
                return  # ЗАВЕРШУЄМО - Excel не потрібен!

            # 3. БД ПОРОЖНЯ - потрібен Excel для міграції
            self.status_bar.showMessage("БД порожня, потрібен Excel для міграції...")
            QApplication.processEvents()

            excel_path = self.config.get("excel_file_path", "D0A02800.xlsx")

            # Якщо шлях порожній, використовуємо дефолтне ім'я файлу
            if not excel_path or excel_path.strip() == "":
                excel_path = "D0A02800.xlsx"

            # Перевірка валідності Excel файлу (ТІЛЬКИ якщо БД порожня)
            is_valid, error_msg = validate_excel_file(excel_path)

            if not is_valid:
                # Excel не знайдено, а БД порожня - пропонуємо вибрати файл
                reply = QMessageBox.question(
                    self,
                    "Потрібен Excel файл",
                    f"База даних порожня, потрібен Excel файл для міграції.\n\n"
                    f"Помилка: {error_msg}\n\n"
                    f"Бажаєте обрати Excel файл зараз?",
                    QMessageBox.Yes | QMessageBox.No
                )

                if reply == QMessageBox.Yes:
                    # Відкриваємо діалог вибору файлу
                    file_path, _ = QFileDialog.getOpenFileName(
                        self,
                        "Оберіть Excel файл з даними",
                        "",
                        "Excel Files (*.xlsx *.xlsm)"
                    )

                    if file_path:
                        # Зберігаємо в конфіг
                        self.config["excel_file_path"] = file_path
                        config_path = os.path.join(os.path.dirname(__file__), '..', '..', 'config', 'settings.json')
                        try:
                            with open(config_path, 'w', encoding='utf-8') as f:
                                json.dump(self.config, f, ensure_ascii=False, indent=2)
                            # Перезавантажуємось з новим файлом
                            self.load_excel_file()
                            return
                        except Exception as e:
                            QMessageBox.critical(self, "Помилка", f"Помилка збереження: {str(e)}")

                self.excel_reader = None
                self.btn_settings.setEnabled(True)
                return

            # Завантажуємо Excel для міграції
            self.status_bar.showMessage("Завантаження Excel файлу для міграції...")
            QApplication.processEvents()

            self.excel_reader = ExcelReader(excel_path)
            self.excel_reader.load_workbook()

            # 4. Пропонуємо виконати міграцію
            reply = QMessageBox.question(
                self,
                "Перша міграція",
                "База даних порожня. Виконати міграцію даних з Excel в БД?\n\n"
                "⚡ ОПТИМІЗОВАНО: міграція тепер займає ~2-3 хвилини\n\n"
                "Після міграції додаток працюватиме значно швидше.",
                QMessageBox.Yes | QMessageBox.No
            )

            if reply == QMessageBox.Yes:
                self._perform_initial_migration()
                # Після міграції - Excel більше не потрібен
                self.excel_reader = None
                self.use_database = True

                # Увімкнюємо кнопки генерації
                self.btn_periods_100.setEnabled(True)
                self.btn_pilgova.setEnabled(True)
                self.btn_import_month.setEnabled(True)
                self.btn_recalculate.setEnabled(True)
                self.btn_add_servicemember.setEnabled(True)
                self.btn_add_period.setEnabled(True)
                self.btn_edit_periods.setEnabled(True)
                self.btn_add_data.setEnabled(True)
                self.btn_settings.setEnabled(True)
                self.status_bar.showMessage("Готово до роботи (джерело: БД)")
            else:
                # Якщо користувач відмовився - працюємо тільки з Excel
                self.use_database = False
                self.btn_periods_100.setEnabled(True)
                self.btn_pilgova.setEnabled(True)
                self.btn_import_month.setEnabled(True)
                self.btn_recalculate.setEnabled(False)  # Для Excel поки вимкнено
                self.btn_add_data.setEnabled(True)
                self.btn_settings.setEnabled(True)

                QMessageBox.information(
                    self,
                    "Інформація",
                    "Додаток працюватиме з Excel файлом.\n"
                    "Для міграції на БД перезапустіть додаток."
                )
                self.status_bar.showMessage(f"Завантажено: {os.path.basename(excel_path)} (джерело: Excel)")

        except Exception as e:
            # Увімкнюємо хоча б налаштування при помилці
            self.btn_settings.setEnabled(True)
            QMessageBox.critical(self, "Помилка", f"Помилка при ініціалізації: {str(e)}")
            self.excel_reader = None
            self.db_manager = None
            self.status_bar.showMessage("Помилка ініціалізації")

    def _perform_initial_migration(self):
        """
        Одноразова міграція Excel → БД з progress dialog
        """
        progress = QProgressDialog("Міграція даних в БД...", None, 0, 100, self)
        progress.setWindowTitle("Міграція")
        progress.setWindowModality(Qt.WindowModal)
        progress.show()

        from PySide6.QtWidgets import QApplication
        QApplication.processEvents()

        try:
            migrator = DataMigration(self.excel_reader, self.db_manager)

            # Виконуємо міграцію
            progress.setLabelText("Міграція аркуша Data...")
            progress.setValue(25)
            QApplication.processEvents()

            stats = migrator.migrate_full_database()

            progress.setValue(100)
            progress.close()

            # Показуємо результат
            QMessageBox.information(
                self,
                "Міграція завершена",
                f"Дані успішно мігровано в БД!\n\n"
                f"Військовослужбовців: {stats['servicemembers']}\n"
                f"Записів: {stats['service_records']}\n"
                f"Періодів 100%: {stats['periods']}\n\n"
                f"Додаток тепер працює швидше!"
            )

            self.use_database = True

        except Exception as e:
            progress.close()
            QMessageBox.critical(
                self,
                "Помилка міграції",
                f"Помилка при міграції даних:\n{str(e)}\n\n"
                f"Додаток працюватиме з Excel файлом."
            )
            self.use_database = False

    def on_periods_100_clicked(self):
        """
        Обробник кнопки "Періоди на 100 тис."
        """
        self.generate_reports(
            sheet_names=["Періоди на 100"],
            template_key="only100",
            report_type="періоди 100 тис."
        )

    def on_pilgova_clicked(self):
        """
        Обробник кнопки "Пільгова вислуга"
        """
        self.generate_reports(
            sheet_names=["Періоди на 100", "Періоди на 30"],
            template_key="pilgova",
            report_type="пільгова вислуга"
        )

    def generate_reports(self, sheet_names, template_key, report_type):
        """
        Генерація рапортів

        Args:
            sheet_names: Список назв аркушів
            template_key: Ключ шаблону в конфігурації
            report_type: Тип рапорту для відображення
        """
        # Перевірка джерела даних
        if self.use_database and not self.db_manager:
            QMessageBox.warning(self, "Попередження", "БД не підключена. Перезапустіть додаток.")
            return
        elif not self.use_database and not self.excel_reader:
            QMessageBox.warning(self, "Попередження", "Excel файл не завантажено. Оберіть файл через Налаштування.")
            return

        # Отримання списку ПІБ (АДАПТОВАНО: з БД або Excel)
        try:
            if self.use_database:
                names = self.db_manager.get_unique_names()
                units = self.db_manager.get_unique_units()
            else:
                names = self.excel_reader.get_unique_names("Data")
                units = self.excel_reader.get_unique_units()
        except Exception as e:
            QMessageBox.critical(self, "Помилка", f"Помилка при читанні даних: {str(e)}")
            return

        if not names:
            QMessageBox.warning(self, "Попередження", "Не знайдено жодного військовослужбовця в файлі.")
            return

        # Відображення діалогу вибору
        dialog = SelectionDialog(names, units, self)
        if dialog.exec():
            mode, value = dialog.get_selection()

            # Визначення списку ПІБ для генерації
            selected_names = []
            if mode == "single":
                selected_names = [value]
            elif mode == "all":
                selected_names = names
            elif mode == "unit":
                # Отримати всі ПІБ з обраного підрозділу
                if self.use_database:
                    # Для БД - фільтруємо з усіх servicemembers
                    all_members = self.db_manager.get_all_servicemembers()
                    selected_names = [sm["name"] for sm in all_members if sm.get("unit") == value]
                else:
                    unit_data = self.excel_reader.get_unit_data(value, "Data")
                    selected_names = list(set([row["name"] for row in unit_data if row.get("name")]))

            if not selected_names:
                QMessageBox.warning(self, "Попередження", "Не обрано жодного військовослужбовця.")
                return

            # Підтвердження
            reply = QMessageBox.question(
                self,
                "Підтвердження",
                f"Створити {len(selected_names)} рапорт(ів) для '{report_type}'?",
                QMessageBox.Yes | QMessageBox.No
            )

            if reply == QMessageBox.Yes:
                self.start_generation(selected_names, sheet_names, template_key, report_type)

    def start_generation(self, names, sheet_names, template_key, report_type):
        """
        Запуск генерації рапортів у фоновому потоці
        """
        # Базова директорія проекту
        # Отримання шляху до шаблону
        template_rel_path = self.config.get("templates", {}).get(template_key, "")
        # Витягуємо ім'я файлу з відносного шляху
        template_filename = os.path.basename(template_rel_path)
        template_path = get_template_path(template_filename)

        if not os.path.exists(template_path):
            QMessageBox.critical(self, "Помилка", f"Шаблон не знайдено: {template_path}")
            return

        # Перевірка на MANUAL маркери в шаблоні
        manual_data = {}
        passport_data_source = None  # Для масової генерації з файлом паспортів

        try:
            doc = Document(template_path)
            generator = ReportGenerator(template_path)
            manual_markers = generator.find_manual_markers(doc)

            if manual_markers:
                # Перевіряємо чи є паспортні маркери
                has_passport_markers = "СЕРІЯ" in manual_markers or "НОМЕР" in manual_markers
                is_mass_generation = len(names) > 1

                # Для масової генерації з паспортними маркерами - показуємо спеціальний діалог
                if has_passport_markers and is_mass_generation:
                    passport_dialog = PassportDataDialog(len(names), self)
                    if passport_dialog.exec():
                        passport_data_source = passport_dialog
                    else:
                        # Користувач скасував - виходимо
                        return

                # Сортуємо маркери: спочатку СЕРІЯ, потім НОМЕР, потім інші по алфавіту
                def marker_sort_key(marker):
                    if marker == "СЕРІЯ":
                        return (0, marker)
                    elif marker == "НОМЕР":
                        return (1, marker)
                    else:
                        return (2, marker)

                manual_markers_sorted = sorted(manual_markers, key=marker_sort_key)

                # Запитуємо користувача для кожного MANUAL маркера
                for marker in manual_markers_sorted:
                    # Пропускаємо паспортні маркери якщо є passport_data_source
                    if passport_data_source and marker in ["СЕРІЯ", "НОМЕР"]:
                        continue

                    prompt = self.get_manual_marker_prompt(marker)
                    # Додаємо підказку що можна пропустити
                    if marker in ["СЕРІЯ", "НОМЕР"]:
                        prompt += "\n(залиште порожнім якщо відсутній)"

                    text, ok = QInputDialog.getText(
                        self,
                        "Введення даних",
                        prompt
                    )

                    if ok:
                        # Якщо користувач натиснув OK, зберігаємо значення (навіть якщо порожнє)
                        value = text.strip() if text else ""
                        # Серія паспорту - завжди великими літерами
                        if marker == "СЕРІЯ" and value:
                            value = value.upper()
                        manual_data[marker] = value
                    else:
                        # Якщо користувач натиснув Cancel - ставимо порожнє значення
                        manual_data[marker] = ""
        except Exception as e:
            QMessageBox.critical(self, "Помилка", f"Помилка при аналізі шаблону: {str(e)}")
            return

        # Отримання директорії виводу
        output_rel_dir = self.config.get("output_directory", "output")
        output_dir = get_output_dir(output_rel_dir)

        # Вимкнення кнопок
        self.btn_periods_100.setEnabled(False)
        self.btn_pilgova.setEnabled(False)
        self.btn_import_month.setEnabled(False)
        self.btn_recalculate.setEnabled(False)
        self.btn_add_data.setEnabled(False)
        self.btn_settings.setEnabled(False)

        # Показати прогрес-бар
        self.progress_bar.setVisible(True)
        self.progress_bar.setValue(0)
        self.progress_bar.setMaximum(len(names))

        # ВИПРАВЛЕНО: Передаємо шлях до БД замість об'єкта (SQLite threading fix)
        db_path = None
        data_source = self.excel_reader

        if self.use_database and self.db_manager:
            # Отримуємо шлях до БД з конфігурації
            db_config = self.config.get("database", {})
            db_path_rel = db_config.get("database_path", "data.db")
            db_path = get_database_path(db_path_rel)
            data_source = None  # Не передаємо об'єкт БД

        # Створення та запуск потоку
        self.thread = ReportGeneratorThread(
            data_source,  # ExcelReader або None
            names,
            sheet_names,
            template_path,
            output_dir,
            manual_data,
            report_type,
            use_database=self.use_database,  # НОВИЙ параметр
            db_path=db_path,  # НОВИЙ параметр - шлях до БД
            passport_data_source=passport_data_source  # Джерело паспортних даних
        )

        self.thread.progress.connect(self.on_progress)
        self.thread.finished.connect(self.on_generation_finished)
        self.thread.error.connect(self.on_generation_error)

        self.thread.start()
        source_msg = "БД" if self.use_database else "Excel"
        self.status_bar.showMessage(f"Генерація рапортів... (джерело: {source_msg})")

    def on_progress(self, current, total):
        """
        Обробник прогресу
        """
        self.progress_bar.setValue(current)
        self.status_bar.showMessage(f"Генерація: {current}/{total}")

    def on_generation_finished(self, success_count, error_count, errors_list=None):
        """
        Обробник завершення генерації
        """
        # Увімкнення кнопок
        self.btn_periods_100.setEnabled(True)
        self.btn_pilgova.setEnabled(True)
        self.btn_import_month.setEnabled(True)
        self.btn_recalculate.setEnabled(True if self.use_database else False)
        self.btn_add_data.setEnabled(True)
        self.btn_settings.setEnabled(True)

        # Сховати прогрес-бар
        self.progress_bar.setVisible(False)

        # Показати результат
        total = success_count + error_count
        message = f"Генерація завершена!\n\nУспішно: {success_count}\nПомилок: {error_count}"

        # Додаємо деталі помилок якщо є
        if error_count > 0 and errors_list:
            errors_text = "\n".join(errors_list[:10])  # Показуємо перші 10 помилок
            if len(errors_list) > 10:
                errors_text += f"\n\n... та ще {len(errors_list) - 10} помилок"
            message += f"\n\nДеталі помилок:\n{errors_text}"

        if error_count == 0:
            QMessageBox.information(self, "Успіх", message)
        else:
            QMessageBox.warning(self, "Завершено з помилками", message)

        self.status_bar.showMessage(f"Готово: {success_count} рапортів створено")

    def on_generation_error(self, error_msg):
        """
        Обробник помилки генерації
        """
        # Увімкнення кнопок
        self.btn_periods_100.setEnabled(True)
        self.btn_pilgova.setEnabled(True)
        self.btn_import_month.setEnabled(True)
        self.btn_recalculate.setEnabled(True if self.use_database else False)
        self.btn_add_data.setEnabled(True)
        self.btn_settings.setEnabled(True)

        # Сховати прогрес-бар
        self.progress_bar.setVisible(False)

        QMessageBox.critical(self, "Помилка", f"Помилка при генерації:\n{error_msg}")
        self.status_bar.showMessage("Помилка при генерації")

    def get_manual_marker_prompt(self, marker: str) -> str:
        """
        Отримати текст запиту для MANUAL маркера

        Args:
            marker: Ключ маркера (наприклад, "СЕРІЯ", "НОМЕР")

        Returns:
            Текст запиту українською
        """
        prompts = {
            "СЕРІЯ": "Введіть серію паспорту:",
            "НОМЕР": "Введіть номер паспорту:",
        }

        # Якщо є специфічний промпт для маркера, використовуємо його
        if marker in prompts:
            return prompts[marker]

        # Інакше генеруємо загальний промпт
        return f"Введіть значення для {marker}:"

    def on_import_month_clicked(self):
        """
        Обробник кнопки "Додати новий місяць"
        ОНОВЛЕНО: Працює з БД або Excel
        """
        # Перевірка джерела даних
        if self.use_database and not self.db_manager:
            QMessageBox.warning(self, "Попередження", "БД не підключена. Перезапустіть додаток.")
            return
        elif not self.use_database and not self.excel_reader:
            QMessageBox.warning(self, "Попередження", "Excel файл не завантажено. Оберіть файл через Налаштування.")
            return

        # Відкриваємо діалог імпорту
        data_source = self.db_manager if self.use_database else self.excel_reader
        dialog = ImportDataDialog(data_source, use_database=self.use_database, parent=self)

        if dialog.exec():
            # Після імпорту
            if self.use_database:
                # Для БД - нічого додаткового робити не треба, дані вже в БД
                QMessageBox.information(
                    self,
                    "Успіх",
                    "Дані успішно додано в базу даних!\n\nМожете одразу генерувати рапорти."
                )
            else:
                # Для Excel - перезавантажуємо файл
                try:
                    self.excel_reader.load_workbook()
                except Exception as e:
                    QMessageBox.critical(self, "Помилка", f"Помилка при перезавантаженні файлу: {str(e)}")

    def on_add_data_clicked(self):
        """
        Обробник кнопки "Імпорт з Excel"
        Опціональний імпорт даних з Excel файлу
        """
        # Запитуємо користувача про файл
        file_path, _ = QFileDialog.getOpenFileName(
            self,
            "Оберіть Excel файл для імпорту",
            "",
            "Excel Files (*.xlsx *.xlsm)"
        )

        if not file_path:
            return

        # Перевіряємо чи БД готова
        if not self.use_database or not self.db_manager:
            QMessageBox.warning(
                self,
                "Попередження",
                "База даних не підключена. Перезапустіть додаток."
            )
            return

        try:
            # Завантажуємо Excel
            excel_reader = ExcelReader(file_path)
            excel_reader.load_workbook()

            # Відкриваємо діалог імпорту
            dialog = ImportDataDialog(self.db_manager, use_database=True, parent=self)
            dialog.excel_reader = excel_reader  # Передаємо для читання даних

            if dialog.exec():
                QMessageBox.information(
                    self,
                    "Успіх",
                    "Дані з Excel успішно імпортовано в базу даних!"
                )

        except Exception as e:
            QMessageBox.critical(
                self,
                "Помилка",
                f"Помилка при імпорті з Excel:\n{str(e)}"
            )

    def on_settings_clicked(self):
        """
        Обробник кнопки "Налаштування"
        """
        file_path, _ = QFileDialog.getOpenFileName(
            self,
            "Оберіть Excel файл",
            "",
            "Excel Files (*.xlsx *.xlsm)"
        )

        if file_path:
            # Оновлення конфігурації
            self.config["excel_file_path"] = file_path

            # Збереження конфігурації
            config_path = get_config_path()

            try:
                with open(config_path, 'w', encoding='utf-8') as f:
                    json.dump(self.config, f, ensure_ascii=False, indent=2)

                # Перезавантаження файлу
                self.load_excel_file()
            except Exception as e:
                QMessageBox.critical(self, "Помилка", f"Помилка при збереженні конфігурації: {str(e)}")

    def on_recalculate_periods_clicked(self):
        """
        Обробник кнопки "Перерахувати періоди"
        Перераховує періоди для всіх військовослужбовців в БД
        """
        if not self.use_database or not self.db_manager:
            QMessageBox.warning(
                self,
                "Попередження",
                "Перерахунок періодів доступний тільки при роботі з базою даних."
            )
            return

        # Підтвердження
        reply = QMessageBox.question(
            self,
            "Підтвердження",
            "Перерахувати періоди для всіх військовослужбовців?\n\n"
            "Це може зайняти кілька хвилин.\n\n"
            "Періоди будуть автоматично оновлені на основі всіх service_records.",
            QMessageBox.Yes | QMessageBox.No
        )

        if reply != QMessageBox.Yes:
            return

        # Отримуємо всіх військовослужбовців
        try:
            all_servicemembers = self.db_manager.get_all_servicemembers()
            total = len(all_servicemembers)

            if total == 0:
                QMessageBox.information(self, "Інформація", "База даних порожня.")
                return

            # Створюємо діалог прогресу
            progress = QProgressDialog("Перерахунок періодів...", "Скасувати", 0, total, self)
            progress.setWindowTitle("Перерахунок періодів")
            progress.setWindowModality(Qt.WindowModal)
            progress.show()

            from PySide6.QtWidgets import QApplication

            # Перераховуємо періоди для кожного
            success_count = 0
            error_count = 0

            for i, servicemember in enumerate(all_servicemembers, 1):
                # Перевіряємо чи користувач скасував
                if progress.wasCanceled():
                    break

                sm_id = servicemember["id"]
                sm_name = servicemember["name"]

                # Оновлюємо прогрес
                progress.setLabelText(f"Обробка {i}/{total}: {sm_name}")
                progress.setValue(i)
                QApplication.processEvents()

                try:
                    self.db_manager.calculate_and_store_periods(sm_id)
                    success_count += 1
                except Exception as e:
                    print(f"[ERROR] Помилка для {sm_name}: {e}")
                    error_count += 1

            # Зберігаємо стан скасування ПЕРЕД закриттям діалогу
            was_canceled = progress.wasCanceled()

            progress.setValue(total)
            progress.close()

            # Показуємо результат
            if was_canceled:
                QMessageBox.information(
                    self,
                    "Скасовано",
                    f"Перерахунок скасовано.\n\nОновлено: {success_count} осіб"
                )
            else:
                message = f"Перерахунок завершено!\n\n"
                message += f"Успішно оновлено: {success_count}\n"
                if error_count > 0:
                    message += f"Помилок: {error_count}\n"
                message += f"\nТепер періоди актуальні для всіх військовослужбовців."

                QMessageBox.information(self, "Успіх", message)

        except Exception as e:
            QMessageBox.critical(
                self,
                "Помилка",
                f"Помилка при перерахунку періодів:\n{str(e)}"
            )

    def on_add_servicemember_clicked(self):
        """
        Обробник кнопки "Додати військовослужбовця"
        """
        if not self.use_database or not self.db_manager:
            QMessageBox.warning(
                self,
                "Попередження",
                "Ця функція доступна тільки при роботі з базою даних."
            )
            return

        dialog = AddServicememberDialog(self.db_manager, self)
        if dialog.exec():
            self.status_bar.showMessage("Військовослужбовця додано")

    def on_add_period_clicked(self):
        """
        Обробник кнопки "Додати період"
        """
        if not self.use_database or not self.db_manager:
            QMessageBox.warning(
                self,
                "Попередження",
                "Ця функція доступна тільки при роботі з базою даних."
            )
            return

        dialog = AddPeriodDialog(self.db_manager, self)
        if dialog.exec():
            self.status_bar.showMessage("Період додано")

    def on_edit_periods_clicked(self):
        """
        Обробник кнопки "Редагувати періоди"
        """
        if not self.use_database or not self.db_manager:
            QMessageBox.warning(
                self,
                "Попередження",
                "Ця функція доступна тільки при роботі з базою даних."
            )
            return

        dialog = EditPeriodsDialog(self.db_manager, self)
        dialog.exec()

    def on_update_clicked(self):
        """
        Обробник кнопки "Оновлення"
        Перевіряє наявність оновлень та пропонує встановити
        """
        from core.updater import (
            check_for_updates, download_update, run_updater,
            get_current_version, open_release_page, GITHUB_REPO
        )
        from PySide6.QtWidgets import QApplication

        self.status_bar.showMessage("Перевірка оновлень...")
        QApplication.processEvents()

        # Перевіряємо оновлення
        release = check_for_updates()

        if release is None:
            current = get_current_version()
            QMessageBox.information(
                self,
                "Оновлення",
                f"Ви використовуєте актуальну версію {current}\n\n"
                f"Оновлень не знайдено."
            )
            self.status_bar.showMessage("Оновлень не знайдено")
            return

        # Є нова версія - питаємо користувача
        current = get_current_version()
        reply = QMessageBox.question(
            self,
            "Доступне оновлення!",
            f"Знайдено нову версію: {release.version}\n"
            f"Поточна версія: {current}\n\n"
            f"{release.description[:200] + '...' if len(release.description) > 200 else release.description}\n\n"
            f"Завантажити та встановити оновлення?",
            QMessageBox.Yes | QMessageBox.No
        )

        if reply != QMessageBox.Yes:
            # Пропонуємо відкрити сторінку релізу
            reply2 = QMessageBox.question(
                self,
                "Відкрити сторінку?",
                "Бажаєте відкрити сторінку релізу в браузері?",
                QMessageBox.Yes | QMessageBox.No
            )
            if reply2 == QMessageBox.Yes:
                open_release_page(release)
            return

        # Завантажуємо оновлення
        progress = QProgressDialog("Завантаження оновлення...", "Скасувати", 0, 100, self)
        progress.setWindowModality(Qt.WindowModal)
        progress.setWindowTitle("Завантаження")
        progress.show()

        def update_progress(current, total):
            if total > 0:
                percent = int(current / total * 100)
                progress.setValue(percent)
                progress.setLabelText(f"Завантаження: {current // 1024} / {total // 1024} KB")
                QApplication.processEvents()

        zip_path = download_update(release, update_progress)
        progress.close()

        if not zip_path:
            QMessageBox.critical(
                self,
                "Помилка",
                "Не вдалося завантажити оновлення.\n\n"
                "Спробуйте завантажити вручну зі сторінки релізу."
            )
            open_release_page(release)
            return

        # Запускаємо updater
        reply = QMessageBox.information(
            self,
            "Готово до оновлення",
            "Оновлення завантажено!\n\n"
            "Програма зараз закриється і запуститься оновлення.\n"
            "Після завершення програма запуститься автоматично.",
            QMessageBox.Ok | QMessageBox.Cancel
        )

        if reply == QMessageBox.Ok:
            if run_updater(zip_path):
                # Закриваємо програму
                QApplication.quit()
            else:
                QMessageBox.critical(
                    self,
                    "Помилка",
                    "Не вдалося запустити оновлення.\n\n"
                    "Спробуйте оновити вручну."
                )
                open_release_page(release)
